#!/usr/bin/env python3
"""Build an Excel test case file from structured JSON using the bundled template."""

from __future__ import annotations

import argparse
import copy
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


HEADERS = [
    "用例名称",
    "所属模块",
    "标签",
    "前置条件",
    "步骤描述",
    "预期结果",
    "编辑模式",
    "备注",
    "用例等级",
]


def text(value: Any, default: str = "") -> str:
    if value is None:
        return default
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def normalize_module(value: Any) -> str:
    module = text(value, "/默认模块")
    if not module:
        module = "/默认模块"
    return module if module.startswith("/") else f"/{module}"


def normalize_priority(value: Any) -> str:
    priority = text(value, "P2").upper()
    return priority if priority in {"P0", "P1", "P2", "P3"} else "P2"


def normalize_tags(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return "｜".join(text(item) for item in value if text(item))
    return text(value)


def numbered(lines: list[str]) -> str:
    return "\n".join(f"[{index}]{line}" for index, line in enumerate(lines, start=1))


def case_steps(case: dict[str, Any]) -> tuple[str, str]:
    steps = case.get("steps") or []
    if isinstance(steps, list) and steps:
        actions: list[str] = []
        expected: list[str] = []
        for item in steps:
            if not isinstance(item, dict):
                item = {"action": item}
            action = text(item.get("action") or item.get("step"))
            result = text(item.get("expected") or item.get("result"))
            if action:
                actions.append(action)
            if result:
                expected.append(result)
        return numbered(actions), numbered(expected)

    description = text(case.get("description") or case.get("step") or case.get("steps"))
    expected = text(case.get("expected") or case.get("result"))
    return description, expected


def flatten_cases(data: dict[str, Any]) -> list[dict[str, Any]]:
    if isinstance(data.get("cases"), list):
        return [case for case in data["cases"] if isinstance(case, dict)]

    flattened: list[dict[str, Any]] = []
    for module in data.get("modules") or []:
        if not isinstance(module, dict):
            continue
        module_name = normalize_module(module.get("name"))
        for case in module.get("cases") or []:
            if not isinstance(case, dict):
                continue
            item = dict(case)
            item.setdefault("module", module_name)
            flattened.append(item)
    return flattened


def copy_row_style(ws, source_row: int, target_row: int) -> None:
    for col in range(1, len(HEADERS) + 1):
        source = ws.cell(source_row, col)
        target = ws.cell(target_row, col)
        if source.has_style:
            target._style = copy.copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy.copy(source.alignment)
        if source.font:
            target.font = copy.copy(source.font)
        if source.fill:
            target.fill = copy.copy(source.fill)
        if source.border:
            target.border = copy.copy(source.border)


def clear_template_rows(ws) -> None:
    if ws.max_row >= 3:
        ws.delete_rows(3, ws.max_row - 2)
    for col in range(1, len(HEADERS) + 1):
        ws.cell(2, col).value = None


def write_cases(input_path: Path, template_path: Path, output_path: Path) -> None:
    data = json.loads(input_path.read_text(encoding="utf-8"))
    cases = flatten_cases(data)

    wb = load_workbook(template_path)
    ws = wb[wb.sheetnames[0]]

    for col, header in enumerate(HEADERS, start=1):
        ws.cell(1, col).value = header

    clear_template_rows(ws)
    for row_index, case in enumerate(cases, start=2):
        if row_index > 2:
            copy_row_style(ws, 2, row_index)

        steps, expected = case_steps(case)
        ws.cell(row_index, 1).value = text(case.get("name"), "未命名用例")
        ws.cell(row_index, 2).value = normalize_module(case.get("module"))
        ws.cell(row_index, 3).value = normalize_tags(case.get("tags"))
        ws.cell(row_index, 4).value = text(case.get("precondition"))
        ws.cell(row_index, 5).value = steps
        ws.cell(row_index, 6).value = expected
        ws.cell(row_index, 7).value = text(case.get("edit_mode"), "STEP") or "STEP"
        ws.cell(row_index, 8).value = text(case.get("remark"))
        ws.cell(row_index, 9).value = normalize_priority(case.get("priority"))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def main() -> None:
    skill_dir = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description="Build an Excel test case file from JSON.")
    parser.add_argument("--input", required=True, type=Path, help="Path to test case JSON.")
    parser.add_argument("--output", required=True, type=Path, help="Output .xlsx path.")
    parser.add_argument(
        "--template",
        type=Path,
        default=skill_dir / "assets" / "excel_case.xlsx",
        help="Template .xlsx path.",
    )
    args = parser.parse_args()

    write_cases(args.input, args.template, args.output)
    print(f"Generated {args.output}")


if __name__ == "__main__":
    main()
